# Copyright (C) 2026 CVAT Custom — Train Group Schedule
# SPDX-License-Identifier: MIT

"""
Pure parser/validation logic for train-group CSV/xlsx/paste uploads.
No Django ORM access; safe to unit-test without a database.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Optional


@dataclass(frozen=True)
class ParsedRow:
    """One validated (train_id, group) pair from an uploaded source."""
    train_id: str
    group: str
    line_no: int  # 1-based; for paste mode without a header, first data row = line 1


@dataclass(frozen=True)
class ValidationError:
    """One row-level rejection reason; safe to serialize back to the client."""
    line: int
    reason: str
    train_id: Optional[str] = None


import csv
import io
import re
from io import BytesIO

_TRAIN_ID_PATTERN = re.compile(r"^[A-Za-z0-9_-]+$")
_MAX_TRAIN_ID_LEN = 100
_MAX_GROUP_LEN = 50
_EXPECTED_HEADER = ("train_id", "group")


def parse_csv(text: str) -> tuple[list[ParsedRow], list[ValidationError]]:
    """
    Parse a CSV string into ParsedRow objects.

    Tolerates: UTF-8 BOM, trailing blank lines, leading/trailing whitespace in cells,
    extra columns beyond train_id,group.

    Rejects: missing/wrong header, empty fields, invalid train_id chars,
    duplicate train_id within the file. All errors collected; partial success is not returned.
    """

    text = text.lstrip("﻿")  # strip BOM if present
    reader = csv.reader(io.StringIO(text))

    rows: list[ParsedRow] = []
    errors: list[ValidationError] = []
    seen_train_ids: dict[str, int] = {}

    try:
        header = next(reader)
    except StopIteration:
        errors.append(ValidationError(line=1, reason="file is empty"))
        return [], errors

    header_normalized = tuple(c.strip().lower() for c in header[:2])
    if header_normalized != _EXPECTED_HEADER:
        errors.append(ValidationError(
            line=1,
            reason=f"expected header 'train_id,group' but got {','.join(header[:2])!r}",
        ))
        return [], errors

    for raw in reader:
        line_no = reader.line_num
        if not raw or all(not c.strip() for c in raw):
            continue  # skip blank line

        train_id = (raw[0] if len(raw) > 0 else "").strip()
        group = (raw[1] if len(raw) > 1 else "").strip()

        row_errors: list[ValidationError] = []

        if not train_id:
            row_errors.append(ValidationError(
                line=line_no, reason="train_id is empty", train_id="",
            ))
        elif len(train_id) > _MAX_TRAIN_ID_LEN:
            row_errors.append(ValidationError(
                line=line_no,
                reason=f"train_id exceeds {_MAX_TRAIN_ID_LEN} characters",
                train_id=train_id,
            ))
        elif not _TRAIN_ID_PATTERN.match(train_id):
            row_errors.append(ValidationError(
                line=line_no,
                reason="train_id contains invalid characters (allowed: A-Z a-z 0-9 _ -)",
                train_id=train_id,
            ))

        if not group:
            row_errors.append(ValidationError(
                line=line_no, reason="group is empty", train_id=train_id,
            ))
        elif len(group) > _MAX_GROUP_LEN:
            row_errors.append(ValidationError(
                line=line_no,
                reason=f"group exceeds {_MAX_GROUP_LEN} characters",
                train_id=train_id,
            ))

        if row_errors:
            errors.extend(row_errors)
            continue

        if train_id in seen_train_ids:
            errors.append(ValidationError(
                line=line_no,
                reason=f"duplicate train_id (first seen on line {seen_train_ids[train_id]})",
                train_id=train_id,
            ))
            continue

        seen_train_ids[train_id] = line_no
        rows.append(ParsedRow(train_id=train_id, group=group, line_no=line_no))

    # If any errors at all, return nothing parsed (all-or-nothing semantics)
    if errors:
        return [], errors
    return rows, []


def parse_xlsx(file_bytes: bytes) -> tuple[list[ParsedRow], list[ValidationError]]:
    """
    Parse the FIRST worksheet of an xlsx file (passed as raw bytes).

    Reads cached values for formulas; if no cached value is present (file never
    opened in Excel), the cell is read as None.

    Normalizes to CSV semantics and then re-uses parse_csv for validation so
    the rules stay identical across input modes.
    """
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        import zipfile
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
        return [], [ValidationError(line=0, reason=f"could not read xlsx: {exc}")]

    ws = wb.worksheets[0]

    # Convert to CSV text — easier than threading line numbers through
    csv_lines: list[str] = []
    for row in ws.iter_rows(values_only=True):
        # Treat all-empty rows as blank
        if not row or all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            csv_lines.append("")
            continue
        cells = []
        for c in row[:2]:  # only first two columns matter
            if c is None:
                cells.append("")
            else:
                cells.append(str(c).strip())
        csv_lines.append(",".join(cells))

    csv_text = "\n".join(csv_lines)
    return parse_csv(csv_text)
