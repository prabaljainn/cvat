from io import BytesIO

from django.test import SimpleTestCase

from cvat.apps.custom.train_group_parser import ParsedRow, ValidationError, parse_csv


class DataClassesTest(SimpleTestCase):
    def test_parsed_row_holds_train_id_and_group(self):
        row = ParsedRow(train_id="3101F", group="A", line_no=2)
        self.assertEqual(row.train_id, "3101F")
        self.assertEqual(row.group, "A")
        self.assertEqual(row.line_no, 2)

    def test_validation_error_holds_line_train_id_reason(self):
        err = ValidationError(line=14, train_id="3101 F", reason="whitespace not allowed")
        self.assertEqual(err.line, 14)
        self.assertEqual(err.reason, "whitespace not allowed")


class ParseCsvHappyPathTest(SimpleTestCase):
    def test_basic_three_rows(self):
        csv = "train_id,group\n3101F,A\n3102F,B\n3103F,A\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(errors, [])
        self.assertEqual(
            rows,
            [
                ParsedRow("3101F", "A", line_no=2),
                ParsedRow("3102F", "B", line_no=3),
                ParsedRow("3103F", "A", line_no=4),
            ],
        )

    def test_tolerates_bom_and_trailing_blank_lines(self):
        csv = "﻿train_id,group\n3101F,A\n\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(errors, [])
        self.assertEqual(rows, [ParsedRow("3101F", "A", line_no=2)])

    def test_trims_whitespace_from_cells(self):
        csv = "train_id,group\n  3101F  ,  A  \n"
        rows, errors = parse_csv(csv)
        self.assertEqual(errors, [])
        self.assertEqual(rows, [ParsedRow("3101F", "A", line_no=2)])

    def test_ignores_unknown_extra_columns(self):
        csv = "train_id,group,notes\n3101F,A,hello\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(errors, [])
        self.assertEqual(rows, [ParsedRow("3101F", "A", line_no=2)])


class ParseCsvValidationTest(SimpleTestCase):
    def test_rejects_missing_header(self):
        csv = "3101F,A\n3102F,B\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(rows, [])
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0].line, 1)
        self.assertIn("header", errors[0].reason.lower())

    def test_rejects_wrong_header(self):
        csv = "trainid,grp\n3101F,A\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(rows, [])
        self.assertEqual(len(errors), 1)
        self.assertIn("header", errors[0].reason.lower())

    def test_rejects_empty_train_id(self):
        csv = "train_id,group\n,A\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(rows, [])
        self.assertEqual(errors, [ValidationError(line=2, reason="train_id is empty", train_id="")])

    def test_rejects_empty_group(self):
        csv = "train_id,group\n3101F,\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(rows, [])
        self.assertEqual(errors, [ValidationError(line=2, reason="group is empty", train_id="3101F")])

    def test_rejects_invalid_train_id_chars(self):
        csv = "train_id,group\n3101 F,A\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(rows, [])
        self.assertEqual(len(errors), 1)
        self.assertIn("invalid characters", errors[0].reason)

    def test_rejects_duplicate_train_id(self):
        csv = "train_id,group\n3101F,A\n3101F,B\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(rows, [])
        self.assertEqual(len(errors), 1)
        self.assertIn("duplicate", errors[0].reason.lower())
        self.assertEqual(errors[0].train_id, "3101F")

    def test_empty_csv_accepted_as_clear(self):
        csv = "train_id,group\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(rows, [])
        self.assertEqual(errors, [])

    def test_collects_multiple_errors(self):
        csv = "train_id,group\n,A\n3102F,\n3101 F,A\n"
        rows, errors = parse_csv(csv)
        self.assertEqual(rows, [])
        self.assertEqual(len(errors), 3)


from openpyxl import Workbook

from cvat.apps.custom.train_group_parser import parse_xlsx


def _xlsx_bytes(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ParseXlsxTest(SimpleTestCase):
    def test_basic_happy_path(self):
        data = _xlsx_bytes([
            ["train_id", "group"],
            ["3101F", "A"],
            ["3102F", "B"],
        ])
        rows, errors = parse_xlsx(data)
        self.assertEqual(errors, [])
        self.assertEqual(
            rows,
            [
                ParsedRow("3101F", "A", line_no=2),
                ParsedRow("3102F", "B", line_no=3),
            ],
        )

    def test_tolerates_trailing_empty_rows(self):
        data = _xlsx_bytes([
            ["train_id", "group"],
            ["3101F", "A"],
            [None, None],
            ["", ""],
        ])
        rows, errors = parse_xlsx(data)
        self.assertEqual(errors, [])
        self.assertEqual(rows, [ParsedRow("3101F", "A", line_no=2)])

    def test_rejects_wrong_header(self):
        data = _xlsx_bytes([
            ["foo", "bar"],
            ["3101F", "A"],
        ])
        rows, errors = parse_xlsx(data)
        self.assertEqual(rows, [])
        self.assertIn("header", errors[0].reason.lower())

    def test_first_sheet_only(self):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Schedule"
        ws1.append(["train_id", "group"])
        ws1.append(["3101F", "A"])
        ws2 = wb.create_sheet("Notes")
        ws2.append(["this should be ignored", "really"])
        buf = BytesIO()
        wb.save(buf)
        rows, errors = parse_xlsx(buf.getvalue())
        self.assertEqual(errors, [])
        self.assertEqual(rows, [ParsedRow("3101F", "A", line_no=2)])

    def test_integer_cell_values_stringified(self):
        # openpyxl returns ints for numeric cells; parser should str() them
        data = _xlsx_bytes([
            ["train_id", "group"],
            [3101, 1],
        ])
        rows, errors = parse_xlsx(data)
        self.assertEqual(errors, [])
        self.assertEqual(rows, [ParsedRow("3101", "1", line_no=2)])

    def test_corrupt_bytes_returns_error(self):
        rows, errors = parse_xlsx(b"not a real xlsx")
        self.assertEqual(rows, [])
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0].line, 0)
